import datetime
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from PIL import Image

from . import reports, services
from .codes import decode_product_code, encode_product_code
from .models import Color, Product, Purchase, PurchaseLine, Sale, SaleLine, Size, Supplier, Variant

MEDIA = tempfile.mkdtemp()


def photo(name='p.png'):
    buf = BytesIO()
    Image.new('RGB', (2400, 1200), 'red').save(buf, 'PNG')
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/png')


def mgmt(prefix, total, initial=0):
    return {f'{prefix}-TOTAL_FORMS': total, f'{prefix}-INITIAL_FORMS': initial,
            f'{prefix}-MIN_NUM_FORMS': 0, f'{prefix}-MAX_NUM_FORMS': 1000}


class CodeTests(TestCase):
    def test_encode_decode(self):
        d = datetime.date(2026, 10, 17)
        self.assertEqual(encode_product_code(d, 1), '4B7A695A9')
        self.assertEqual(decode_product_code('4B7A695A9'), (d, 1))
        self.assertEqual(decode_product_code(encode_product_code(d, 999)), (d, 999))


@override_settings(MEDIA_ROOT=MEDIA)
class FlowTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.boss = User.objects.create_superuser('boss', password='x')
        self.client.force_login(self.boss)
        self.sup = Supplier.objects.create(code='S01', name='五分埔A')
        self.bk, self.wh = Color.objects.get(code='BK'), Color.objects.get(code='WH')
        self.m, self.l = Size.objects.get(code='M'), Size.objects.get(code='L')

    def make_product(self):
        resp = self.client.post('/admin/stock/product/add/', {
            'name': '素T', 'category': 'top', 'supplier': self.sup.pk, 'supplier_code': 'A-77',
            'list_price': '390', 'image': photo(), 'note': '',
            'add_colors': [self.bk.pk, self.wh.pk], 'add_sizes': [self.m.pk, self.l.pk],
            **mgmt('variants', 0),
        })
        self.assertEqual(resp.status_code, 302, resp.content.decode()[:3000])
        return Product.objects.get(name='素T')

    def purchase(self, variant, qty, price, date=None, with_photo=True):
        data = {'date': (date or datetime.date.today()).isoformat(), 'supplier': self.sup.pk, 'note': '',
                **mgmt('photos', 1), **mgmt('lines', 1),
                'lines-0-variant': variant.pk, 'lines-0-qty': qty, 'lines-0-unit_price': price}
        if with_photo:
            data['photos-0-image'] = photo()
        return self.client.post('/admin/stock/purchase/add/', data)

    def sell(self, variant, qty, price=''):
        return self.client.post('/admin/stock/sale/add/', {
            'date': datetime.date.today().isoformat(), 'channel': 'store', 'payment': 'cash', 'discount': '0',
            'note': '', **mgmt('lines', 1),
            'lines-0-variant': variant.pk, 'lines-0-qty': qty, 'lines-0-unit_price': price})

    def test_full_flow(self):
        p = self.make_product()
        code_date, serial = decode_product_code(p.code)
        self.assertEqual((code_date, serial), (p.code_date, 1))
        self.assertEqual(p.variants.count(), 4)
        self.assertLessEqual(max(Image.open(p.image.path).size), 1600)  # 已壓縮
        v = Variant.objects.get(product=p, color=self.bk, size=self.m)
        self.assertEqual(v.sku, f'{p.code}-BK-M')

        # 沒附照片不能存
        resp = self.purchase(v, 10, '100', with_photo=False)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '至少要附一張照片')
        self.assertEqual(Purchase.objects.count(), 0)

        # 移動加權平均：10@100 + 10@130 → 115
        self.assertEqual(self.purchase(v, 10, '100').status_code, 302)
        self.assertEqual(self.purchase(v, 10, '130').status_code, 302)
        v.refresh_from_db()
        self.assertEqual((v.qty_on_hand, v.avg_cost), (20, Decimal('115.00')))
        self.assertTrue(Purchase.objects.first().no.startswith('PO'))

        # 庫存不足擋下
        resp = self.sell(v, 21)
        self.assertContains(resp, '庫存不足')
        # 單價留空 → 建議售價
        self.assertEqual(self.sell(v, 5).status_code, 302)
        line = SaleLine.objects.get()
        self.assertEqual(line.unit_price, Decimal('390'))
        self.assertEqual(line.move.unit_cost, Decimal('115.00'))
        v.refresh_from_db()
        self.assertEqual(v.qty_on_hand, 15)

        # 作廢銷貨 → 庫存回補
        services.void(Sale.objects.get())
        v.refresh_from_db()
        self.assertEqual(v.qty_on_hand, 20)

        # 補登過去日期的進貨，成本依日期重算
        self.assertEqual(self.sell(v, 2).status_code, 302)
        self.purchase(v, 20, '40', date=datetime.date.today() - datetime.timedelta(days=30))
        v.refresh_from_db()
        # 40×20 + 100×10 + 130×10 = 3100 / 40 = 77.5
        self.assertEqual(SaleLine.objects.filter(sale__status='posted').get().move.unit_cost, Decimal('77.50'))
        self.assertEqual(v.qty_on_hand, 38)

        # 報表與 Excel
        for name in reports.REPORTS:
            self.assertEqual(self.client.get(f'/reports/{name}/').status_code, 200, name)
            resp = self.client.get(f'/reports/{name}/?format=xlsx')
            self.assertEqual(resp.status_code, 200, name)
            self.assertTrue(resp.content.startswith(b'PK'))
        q = reports.quarterly(reports_params()).rows[-1]
        self.assertEqual(q[2], Decimal('780'))     # 銷貨總額 2×390
        self.assertEqual(q[3], Decimal('155.00'))  # 銷貨成本 2×77.5

    def test_lookup(self):
        p = self.make_product()
        v = Variant.objects.get(product=p, color=self.bk, size=self.m)
        self.purchase(v, 7, '100')
        for q in (p.code, p.code.lower(), v.sku, '素'):
            resp = self.client.get('/lookup/', {'q': q})
            self.assertContains(resp, '總庫存：7', msg_prefix=q)
        self.assertContains(self.client.get('/lookup/', {'q': 'ZZZ'}), '找不到')

    def test_three_reports_file(self):
        from openpyxl import load_workbook
        p = self.make_product()
        v = Variant.objects.get(product=p, color=self.bk, size=self.m)
        self.purchase(v, 10, '100')
        self.sell(v, 3, '300')
        with override_settings(REPORT_DIR=Path(MEDIA) / '報表'):
            resp = self.client.post('/update-reports/')
            self.assertEqual(resp.status_code, 302)
            wb = load_workbook(Path(MEDIA) / '報表' / '庫存系統報表.xlsx')
            self.assertEqual(wb.sheetnames, ['進貨系統', '銷貨系統', '收益計算'])
            buy = [c.value for c in wb['進貨系統'][2]]
            self.assertEqual(buy[2:10], ['五分埔A', '素T', '黑/M', 10, 100, 1000, 'A-77', p.code])
            self.assertEqual([c.value for c in wb['銷貨系統'][2]][6:9], [3, 300, 900])
            self.assertEqual([c.value for c in wb['收益計算'][2]][1:5], [1000, 900, 300, 600])
            # 再賣一筆後更新，報表要跟著變
            self.sell(v, 1, '300')
            self.client.post('/update-reports/')
            wb = load_workbook(Path(MEDIA) / '報表' / '庫存系統報表.xlsx')
            self.assertEqual(wb['銷貨系統'].max_row, 3)
            self.assertEqual(wb['收益計算']['C2'].value, 1200)
            self.assertEqual(self.client.get('/update-reports/?download=1').status_code, 200)

    def test_clerk_cannot_see_cost(self):
        p = self.make_product()
        clerk = User.objects.create_user('clerk', password='x', is_staff=True)
        clerk.groups.add(Group.objects.get(name='店員'))
        self.client.force_login(clerk)
        resp = self.client.get('/reports/inventory/')
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, '平均成本')
        self.assertEqual(self.client.get('/reports/quarterly/').status_code, 403)
        self.assertEqual(self.client.get('/admin/stock/purchase/add/').status_code, 403)
        self.assertEqual(self.client.get('/admin/stock/sale/add/').status_code, 200)
        self.assertContains(self.client.get('/lookup/', {'q': p.code}), '總庫存')
        self.assertEqual(self.client.get('/update-reports/').status_code, 403)


def reports_params():
    today = datetime.date.today()
    return {'today': today, 'start': today.replace(day=1), 'end': today, 'days': 60}
