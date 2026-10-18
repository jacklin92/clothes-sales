import datetime
from decimal import Decimal
from urllib.parse import quote

from django.conf import settings
from django.contrib import admin, messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import reports
from .images import thumbnail_bytes
from .reports import REPORTS, three_reports

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
THREE_REPORTS_FILE = '庫存系統報表.xlsx'

MENU = [
    ('inventory', '即時庫存表'), ('purchases', '進貨報表'), ('sales', '銷貨報表'),
    ('quarterly', '季度收益'), ('slow', '滯銷品'), ('top', '熱銷排行'),
]


def _params(request):
    today = timezone.localdate()

    def date(name, default):
        try:
            return datetime.date.fromisoformat(request.GET.get(name, ''))
        except ValueError:
            return default

    try:
        days = max(1, int(request.GET.get('days', 60)))
    except ValueError:
        days = 60
    return {'today': today, 'start': date('start', today.replace(day=1)), 'end': date('end', today), 'days': days}


@staff_member_required
def report(request, name):
    if name not in REPORTS:
        raise Http404
    params = _params(request)
    rep = REPORTS[name](params)
    can_cost = request.user.has_perm('stock.view_cost')
    if rep.needs_cost and not can_cost:
        raise PermissionDenied
    if not can_cost:
        rep.strip_cost()
    if request.GET.get('format') == 'xlsx':
        return _xlsx(rep)
    return render(request, 'stock/report.html', {
        **admin.site.each_context(request), 'title': rep.title, 'rep': rep, 'name': name, 'menu': MENU,
        'params': params, 'headers': rep.clean_headers(), 'image_cols': rep.image_cols(),
    })


def _xlsx(rep):
    wb = Workbook()
    _write_sheet(wb.active, rep, rep.title)
    resp = HttpResponse(content_type=XLSX)
    fname = f'{rep.title.split(" ")[0]}_{timezone.localdate():%Y%m%d}.xlsx'
    resp['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(fname)}"
    wb.save(resp)
    return resp


def write_three_reports(path):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, rep in three_reports(timezone.localdate()):
        _write_sheet(wb.create_sheet(), rep, sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


@staff_member_required
def update_reports(request):
    if not request.user.has_perm('stock.view_cost'):
        raise PermissionDenied
    path = settings.REPORT_DIR / THREE_REPORTS_FILE
    if request.GET.get('download') and path.exists():
        return FileResponse(open(path, 'rb'), as_attachment=True, filename=path.name)
    if request.method == 'POST':
        try:
            write_three_reports(path)
            messages.success(request, f'已更新：{path}')
        except PermissionError:
            messages.error(request, f'無法寫入 {path.name}，請先關閉 Excel 中開啟的這個檔案再試一次。')
        return redirect('update_reports')
    updated = datetime.datetime.fromtimestamp(path.stat().st_mtime) if path.exists() else None
    return render(request, 'stock/update_reports.html', {
        **admin.site.each_context(request), 'title': '更新三大報表', 'path': path, 'updated': updated,
    })


@staff_member_required
def lookup(request):
    if not request.user.has_perm('stock.view_variant'):
        raise PermissionDenied
    q = request.GET.get('q', '')
    products = reports.lookup(q)
    results = [(p, *reports.stock_matrix(p)) for p in products] if len(products) == 1 else []
    return render(request, 'stock/lookup.html', {
        **admin.site.each_context(request), 'title': '貨號查詢庫存', 'q': q, 'products': products,
        'results': results,
    })


def _write_sheet(ws, rep, title):
    ws.title = title[:31]
    ws.append(rep.clean_headers())
    for c in ws[1]:
        c.font = Font(bold=True)
    img_cols = rep.image_cols()
    for r, row in enumerate(rep.rows, start=2):
        for c, val in enumerate(row, start=1):
            if c - 1 in img_cols:
                if val:
                    try:
                        pic = XLImage(thumbnail_bytes(val.path))
                    except OSError:
                        continue
                    ws.add_image(pic, f'{get_column_letter(c)}{r}')
                    ws.row_dimensions[r].height = 62
                continue
            cell = ws.cell(r, c, float(val) if isinstance(val, Decimal) else val)
            if isinstance(val, datetime.date):
                cell.number_format = 'yyyy-mm-dd'
    for c, h in enumerate(rep.clean_headers(), start=1):
        ws.column_dimensions[get_column_letter(c)].width = 12 if c - 1 in img_cols else max(10, len(h) * 2 + 4)
